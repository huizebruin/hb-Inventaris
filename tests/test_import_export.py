import io, json
from openpyxl import Workbook, load_workbook


def upload(c, name, data, **form):
    return c.post('/api/import/preview', data={'file': (io.BytesIO(data), name), **form}, content_type='multipart/form-data')


CSV = ("naam;IPN;voorraad;categorie;locatie\n"
       "Weerstand 1k;R-1K;10;Weerstanden (THT);Lade 1\n"
       "Condensator;C-100;5;Onbekende cat;Doos A\n"
       ";X-1;3;;\n"
       "Kapot;K-1;abc;;\n"
       "Weerstand 1k;R-1K;1;;\n"
       "ESP32 DevKit;;9;;\n").encode()


def count(client): return len(client.get('/api/parts').get_json())


def test_csv_preview_writes_nothing_and_reports(client):
    before = count(client)
    j = upload(client, 'x.csv', CSV).get_json()
    c = j['counts']
    assert count(client) == before                      # preview wijzigt niets
    assert c['total'] == 6 and c['invalid'] == 2 and c['missing_required'] == 1
    assert c['duplicate_in_file'] == 1 and c['duplicate_in_db'] == 1   # ESP32 DevKit bestaat al (demo)
    assert c['create'] == 2 and c['skip'] == 1
    assert {i['row'] for i in j['invalid_rows']} == {4, 5}
    assert any('Onbekende cat' in w['message'] for w in j['warnings'])
    assert j['mapping']['name'] == 'naam' and j['missing_required'] == []


def test_cancel_prevents_commit(client):
    token = upload(client, 'x.csv', CSV).get_json()['token']
    assert client.delete(f'/api/import/{token}').get_json()['cancelled']
    assert client.post('/api/import/commit', json={'token': token, 'confirm': True}).status_code == 404
    assert count(client) == 3


def test_commit_needs_confirm_then_applies(client):
    token = upload(client, 'x.csv', CSV).get_json()['token']
    assert client.post('/api/import/commit', json={'token': token}).status_code == 400
    r = client.post('/api/import/commit', json={'token': token, 'confirm': True}).get_json()
    assert r['created'] == 2 and r['skipped'] == 2 and r['invalid'] == 2 and r['safety_backup'].endswith('pre-import.db')
    names = {p['name']: p for p in client.get('/api/parts').get_json()}
    assert names['Weerstand 1k']['catId'] == 202 and names['Weerstand 1k']['locId'] == 1
    assert names['Condensator']['catId'] is None
    assert client.post('/api/parts', json={'name': 'Daarna'}).status_code == 201   # tellers kloppen


def test_update_mode_changes_stock_with_history(client):
    csv_ = b"naam;IPN;voorraad\nESP32 DevKit;DEMO-002;42\n"
    j = upload(client, 'u.csv', csv_, duplicate_mode='update').get_json()
    assert j['counts']['update'] == 1 and j['updates'][0]['diff']['stock'] == [3, 42]
    client.post('/api/import/commit', json={'token': j['token'], 'confirm': True})
    assert client.get('/api/parts').get_json()[1]['stock'] == 42
    assert any(h['delta'] == 39 for h in client.get('/api/history?partId=2').get_json())


def test_remap_columns_via_preview(client):
    csv_ = b"a;b\nTest;5\n"
    j = upload(client, 'm.csv', csv_).get_json()
    assert j['missing_required'] == ['name']
    assert client.post('/api/import/commit', json={'token': j['token'], 'confirm': True}).status_code == 400
    j2 = client.post('/api/import/preview', json={'token': j['token'], 'mapping': {'name': 'a', 'stock': 'b'}}).get_json()
    assert j2['counts']['create'] == 1 and j2['missing_required'] == []


def test_failed_import_leaves_db_untouched(srv, client, monkeypatch):
    token = upload(client, 'x.csv', CSV).get_json()['token']
    monkeypatch.setattr(srv, '_apply_part_insert', lambda *a, **k: (_ for _ in ()).throw(RuntimeError('boem')))
    r = client.post('/api/import/commit', json={'token': token, 'confirm': True})
    assert r.status_code == 500 and count(client) == 3


def test_bad_files_rejected(client):
    assert upload(client, 'x.csv', b'naam\n').status_code == 400          # alleen kop
    assert upload(client, 'x.exe', b'abc').status_code == 400
    assert upload(client, 'x.json', b'{kapot').status_code == 400
    assert upload(client, 'x.xlsx', b'geen zip').status_code == 400


def test_excel_roundtrip(client):
    r = client.get('/api/export/xlsx')
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data)); rows = list(wb['Onderdelen'].values)
    assert rows[0][1] == 'name' and len(rows) == 4
    j = upload(client, 'e.xlsx', r.data, duplicate_mode='update').get_json()
    assert j['counts']['update'] == 3 and j['counts']['invalid'] == 0
    wb = Workbook(); ws = wb.active; ws.append(['Naam', 'Voorraad']); ws.append(['Uit Excel', 7]); ws.append(['Fout', -1])
    bio = io.BytesIO(); wb.save(bio)
    j = upload(client, 'n.xlsx', bio.getvalue()).get_json()
    assert j['counts']['create'] == 1 and j['counts']['invalid'] == 1


def test_csv_export_neutralizes_formulas_and_roundtrips(client):
    client.post('/api/parts', json={'name': '=HYPERLINK("x")', 'IPN': 'F-1'})
    txt = client.get('/api/export/csv').data.decode('utf-8-sig')
    assert "'=HYPERLINK" in txt
    j = upload(client, 'r.csv', txt.encode('utf-8-sig'), duplicate_mode='update').get_json()
    assert j['counts']['invalid'] == 0 and j['counts']['update'] == 4
    assert not any('name' in u['diff'] for u in j['updates'])


def test_json_export_selection_and_merge(client):
    full = client.get('/api/export/json').get_json()
    assert full['scope'] == 'full' and len(full['parts']) == 3 and full['categories'] and 'projects' in full
    sel = client.get('/api/export/json?ids=1,2').get_json()
    assert sel['scope'] == 'selection' and [p['id'] for p in sel['parts']] == [1, 2] and 'projects' not in sel
    j = upload(client, 's.json', json.dumps(sel).encode()).get_json()
    assert j['counts']['duplicate_in_db'] == 2 and not j['replace_possible']
    # een selectie mag nooit als volledige restore dienen
    r = client.post('/api/import/commit', json={'token': j['token'], 'confirm': True, 'options': {'mode': 'replace'}})
    assert r.status_code == 422


def test_json_full_restore_replaces_everything(client):
    full = client.get('/api/export/json?history=1').get_json()
    client.delete('/api/parts/1'); client.post('/api/parts', json={'name': 'Weg ermee'})
    j = upload(client, 'f.json', json.dumps(full).encode(), mode='replace').get_json()
    assert j['replace_possible'] and j['replace_summary']['incoming']['parts'] == 3
    r = client.post('/api/import/commit', json={'token': j['token'], 'confirm': True}).get_json()
    assert r['ok'] and r['mode'] == 'replace'
    assert sorted(p['name'] for p in client.get('/api/parts').get_json()) == sorted(p['name'] for p in full['parts'])
    assert client.post('/api/parts', json={'name': 'Nieuw'}).status_code == 201


def test_json_validation_errors(client):
    bad = {'parts': [{'name': 'a', 'tags': 'geen lijst'}]}
    j = upload(client, 'b.json', json.dumps(bad).encode()).get_json()
    assert j['json_errors'] and not j['replace_possible']
